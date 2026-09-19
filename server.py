#!/usr/bin/env python3
"""
Private registration server for The Scholarship Register.

Run:
    python3 server.py

Then open:
    http://127.0.0.1:8000/

The Excel workbook is stored beside this file but is NEVER served by this
web server. Visitors only receive index.html and API responses.
"""
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
import json
import threading
from datetime import datetime, timezone
import re

try:
    import openpyxl
except ImportError:
    raise SystemExit("Missing dependency. Run: python3 -m pip install -r requirements.txt")

BASE = Path(__file__).resolve().parent
HTML = BASE / "index.html"
WORKBOOK = BASE / "registrations.xlsx"
HOST = "127.0.0.1"
PORT = 8000

lock = threading.Lock()
HEADERS = [
    "Signed up at", "Name", "Age", "Phone", "Address",
    "Permanent address", "LinkedIn", "Certificate file name"
]

def ensure_workbook():
    if WORKBOOK.exists():
        wb = openpyxl.load_workbook(WORKBOOK)
        if "Registrations" not in wb.sheetnames:
            ws = wb.create_sheet("Registrations")
        else:
            ws = wb["Registrations"]
        if ws.max_row == 0 or all(v is None for v in ws[1]):
            ws.append(HEADERS)
        elif [ws.cell(1, c).value for c in range(1, len(HEADERS)+1)] != HEADERS:
            # Keep the user's existing workbook data intact; only add missing headers
            # if the first row is empty.
            pass
        wb.save(WORKBOOK)
        wb.close()
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Registrations"
        ws.append(HEADERS)
        wb.save(WORKBOOK)
        wb.close()

def clean(value, max_len=2000):
    if value is None:
        return ""
    value = str(value).strip()
    # Prevent Excel formula injection when a visitor enters =, +, -, or @.
    if value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value[:max_len]

def valid_phone(phone):
    return bool(re.fullmatch(r"\d{10}", phone))

def save_registration(data):
    name = clean(data.get("name"), 200)
    age = data.get("age")
    phone = re.sub(r"\D", "", str(data.get("phone", "")))[-10:]
    address = clean(data.get("address"), 2000)
    home = clean(data.get("homeAddress"), 2000)
    linkedin = clean(data.get("linkedin"), 500)
    certificate = clean(data.get("certificateName"), 255)

    if not name or not valid_phone(phone):
        raise ValueError("Name and a valid 10-digit mobile number are required.")
    try:
        age = int(age)
    except (TypeError, ValueError):
        raise ValueError("Age must be a number.")
    if age < 1 or age > 100:
        raise ValueError("Age must be between 1 and 100.")
    if not address or not home:
        raise ValueError("Both addresses are required.")

    with lock:
        ensure_workbook()
        wb = openpyxl.load_workbook(WORKBOOK)
        ws = wb["Registrations"]

        # Server-side duplicate check.
        for row in ws.iter_rows(min_row=2, values_only=True):
            existing = re.sub(r"\D", "", str(row[3] or ""))[-10:]
            if existing == phone:
                wb.close()
                raise ValueError("This mobile number is already registered.")

        timestamp = data.get("signedUpAt") or datetime.now(timezone.utc).isoformat()
        ws.append([
            clean(timestamp, 80),
            name,
            age,
            phone,
            address,
            home,
            linkedin,
            certificate
        ])
        wb.save(WORKBOOK)
        wb.close()

class Handler(BaseHTTPRequestHandler):
    server_version = "ScholarshipRegisterPrivate/1.0"

    def send_json(self, status, payload):
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        path = urlparse(self.path).path
        if path in ("/", "/index.html"):
            try:
                body = HTML.read_bytes()
            except FileNotFoundError:
                return self.send_json(500, {"success": False, "message": "index.html is missing."})
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(body)
            return
        if path == "/api/health":
            return self.send_json(200, {"success": True})
        # Deliberately do not expose registrations.xlsx or directory contents.
        self.send_json(404, {"success": False, "message": "Not found."})

    def do_POST(self):
        path = urlparse(self.path).path
        if path != "/api/register":
            return self.send_json(404, {"success": False, "message": "Not found."})

        length = int(self.headers.get("Content-Length", "0"))
        if length > 100_000:
            return self.send_json(413, {"success": False, "message": "Request too large."})

        try:
            raw = self.rfile.read(length)
            data = json.loads(raw.decode("utf-8"))
            save_registration(data)
        except ValueError as e:
            return self.send_json(400, {"success": False, "message": str(e)})
        except Exception as e:
            print("Registration save error:", repr(e))
            return self.send_json(500, {"success": False, "message": "Private registration storage failed."})

        self.send_json(200, {"success": True, "message": "Registration saved."})

    def log_message(self, fmt, *args):
        # Keep visitor registration details out of the terminal log.
        if self.path.startswith("/api/register"):
            print("Registration request received.")
        else:
            super().log_message(fmt, *args)

if __name__ == "__main__":
    ensure_workbook()
    print(f"Private registration server running at http://{HOST}:{PORT}/")
    print(f"Private workbook: {WORKBOOK.name} (not web-accessible)")
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()
